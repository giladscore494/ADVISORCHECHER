"""Fetch a URL and extract readable text from HTML, PDF, JSON (incl. CKAN APIs), CSV or XLSX."""

import csv
import io
import ipaddress
import json
import re
import socket
import time
import zipfile
from dataclasses import asdict, dataclass, field
from typing import Any
from urllib.parse import urljoin, urlparse
from urllib.robotparser import RobotFileParser

import requests
from bs4 import BeautifulSoup

import documents
from config import get_int
from evidence import strip_controls

USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/124.0 Safari/537.36 RegulatoryOpportunityHunter/0.1"
)
TIMEOUT = (10, 30)  # connect, read
# Legal PDFs (laws with all their schedules) can be large; a download over the cap fails loudly, never silently.
MAX_DOWNLOAD_BYTES = get_int("DOCUMENT_MAX_BYTES", 40 * 1024 * 1024)
# Kept for callers that ask for a bounded text; fetch_url itself no longer truncates (max_chars=None).
MAX_TEXT_CHARS = 15000
MAX_REDIRECTS = 5
MAX_TABLE_ROWS = 200  # rows rendered per CSV file / XLSX sheet
MAX_XLSX_SHEETS = 5
MAX_XLSX_UNCOMPRESSED = 100 * 1024 * 1024  # zip-bomb guard
MAX_CELL_CHARS = 200

# Transient-failure retries (429 / 5xx / connection errors) with exponential backoff and a strict cap.
MAX_RETRIES = 2
BACKOFF_BASE_S = 1.0
MAX_BACKOFF_S = 10.0
TRANSIENT_STATUSES = {429, 500, 502, 503, 504}
ROBOTS_AGENT = "RegulatoryOpportunityHunter"
MAX_ROBOTS_BYTES = 512 * 1024
_sleep = time.sleep  # patched in tests
_robots_cache: dict[str, RobotFileParser | None] = {}

XLSX_TYPES = ("spreadsheetml", "ms-excel", "officedocument.spreadsheet")
CKAN_RESOURCE_FIELDS = (
    "name", "description", "format", "mimetype", "url", "package_id", "id",
    "created", "last_modified", "metadata_modified", "size", "datastore_active",
)

STRIP_TAGS = ["script", "style", "noscript", "svg", "nav", "header", "footer", "aside", "form", "iframe", "button"]


@dataclass
class FetchResult:
    url: str
    final_url: str
    ok: bool
    source_type: str  # html | pdf | json | ckan_resource | ckan_dataset | ckan_datastore | csv | xlsx | text | unknown
    title: str = ""
    text: str = ""
    error: str = ""
    http_status: int | None = None
    truncated: bool = False
    # Downloadable file behind a CKAN resource (data.gov.il resource_show); fetch it to read the data.
    resource_url: str = ""
    metadata: dict = field(default_factory=dict)
    attempts: int = 1
    # True for 401/403 or robots.txt disallow: access is restricted and must not be bypassed.
    access_restricted: bool = False
    # Page-by-page extraction (documents.Extracted) and the downloaded bytes, for the document store.
    document: Any = None
    raw: bytes = b""

    def to_dict(self) -> dict:
        return {k: v for k, v in asdict(self).items() if k not in ("document", "raw")}


def _is_public_host(host: str) -> bool:
    """Reject localhost / private network targets (the model chooses URLs)."""
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror:
        return False
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            return False
    return True


def _check_url(url: str) -> str | None:
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        return "Only http(s) URLs are supported."
    if not _is_public_host(parsed.hostname):
        return "Host is not publicly resolvable or is a private address."
    return None


def _truncate(text: str, limit: int) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit] + "\n[... truncated ...]", True


def extract_html_text(html: str | bytes) -> tuple[str, str]:
    """Return (title, readable_text) from an HTML document."""
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else ""
    for tag in soup(STRIP_TAGS):
        tag.decompose()
    root = soup.find("main") or soup.find("article") or soup.body or soup
    text = root.get_text("\n", strip=True)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return title, text.strip()


def extract_pdf_text(data: bytes) -> tuple[str, str]:
    """Return (title, text) with every page of a PDF ("[page n]" markers). Raises ValueError if no text."""
    doc = documents.extract_pdf(data)
    return doc.title, pages_text(doc)


def pages_text(doc) -> str:
    return "\n\n".join(f"[{doc.unit} {i}]\n{t}" for i, t in enumerate(doc.pages, start=1) if t)


def _cell(value) -> str:
    text = "" if value is None else str(value).strip().replace("\n", " ")
    return text[:MAX_CELL_CHARS]


def _render_rows(header: list, rows: list[list], total: int | None) -> str:
    lines = ["Columns: " + " | ".join(_cell(h) for h in header)] if header else []
    lines += [" | ".join(_cell(c) for c in row) for row in rows]
    shown = len(rows)
    if total is not None and total > shown:
        lines.append(f"[showing first {shown} of {total} data rows]")
    return "\n".join(lines)


def decode_text(data: bytes) -> str:
    """Decode bytes as UTF-8, then Windows-1255 (common in Israeli government files), then Latin-1."""
    for encoding in ("utf-8-sig", "cp1255"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def iter_csv_rows(data: bytes):
    """Return (header, row_iterator) for CSV bytes. Rows are lists of strings; blank rows skipped."""
    text = decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    header = next(reader, [])
    return header, (row for row in reader if any(c.strip() for c in row))


def extract_csv_text(data: bytes) -> str:
    header, rows_iter = iter_csv_rows(data)
    rows, total = [], 0
    for row in rows_iter:
        total += 1
        if len(rows) < MAX_TABLE_ROWS:
            rows.append(row)
    if not header and not rows:
        raise ValueError("CSV file is empty.")
    return _render_rows(header, rows, total)


def open_xlsx(data: bytes):
    """Open an XLSX workbook read-only after a zip-bomb check. Raises ValueError."""
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            if sum(i.file_size for i in zf.infolist()) > MAX_XLSX_UNCOMPRESSED:
                raise ValueError("XLSX file is too large when uncompressed.")
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except zipfile.BadZipFile as exc:
        raise ValueError("Not a valid XLSX file (legacy .xls is not supported).") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"XLSX could not be parsed: {exc}") from exc


def iter_xlsx_sheets(wb):
    """Yield (sheet_title, header, row_iterator, total_or_None) for the first MAX_XLSX_SHEETS sheets."""
    for ws in wb.worksheets[:MAX_XLSX_SHEETS]:
        rows_iter = ws.iter_rows(values_only=True)
        header = list(next(rows_iter, None) or [])
        rows = (list(r) for r in rows_iter if any(c not in (None, "") for c in r))
        yield ws.title, header, rows, (ws.max_row - 1) if ws.max_row else None


def extract_xlsx_text(data: bytes) -> str:
    wb = open_xlsx(data)
    parts = []
    try:
        for title, header, rows_iter, total in iter_xlsx_sheets(wb):
            rows = []
            for row in rows_iter:
                if len(rows) >= MAX_TABLE_ROWS:
                    break
                rows.append(row)
            parts.append(f"[sheet: {title}]\n" + _render_rows(header, rows, total))
        if len(wb.worksheets) > MAX_XLSX_SHEETS:
            parts.append(f"[{len(wb.worksheets) - MAX_XLSX_SHEETS} more sheets not shown]")
    finally:
        wb.close()
    text = "\n\n".join(parts).strip()
    if not text:
        raise ValueError("XLSX contains no readable cells.")
    return text


def json_records(doc) -> list[dict] | None:
    """Find a list of records in common JSON layouts (list of objects, CKAN/GeoJSON wrappers)."""
    if isinstance(doc, list) and doc and all(isinstance(r, dict) for r in doc[:20]):
        return doc
    if isinstance(doc, dict):
        if isinstance(doc.get("features"), list):  # GeoJSON
            return [f.get("properties") or {} for f in doc["features"] if isinstance(f, dict)]
        for key in ("records", "result", "data", "items", "rows"):
            value = doc.get(key)
            if isinstance(value, dict):
                nested = json_records(value)
                if nested is not None:
                    return nested
            elif isinstance(value, list) and value and all(isinstance(r, dict) for r in value[:20]):
                return value
    return None


def detect_kind(content_type: str, final_url: str, data: bytes) -> str:
    """Classify a download as pdf | json | xlsx | csv | html."""
    content_type = (content_type or "").lower()
    path = urlparse(final_url).path.lower()
    head = data.lstrip()[:1]
    if "pdf" in content_type or data[:5] == b"%PDF-" or path.endswith(".pdf"):
        return "pdf"
    if "json" in content_type or (head in (b"{", b"[") and "html" not in content_type):
        return "json"
    # CSV downloads are often labelled application/vnd.ms-excel, so a .csv path wins.
    if path.endswith((".xlsx", ".xlsm")) or (any(t in content_type for t in XLSX_TYPES) and not path.endswith(".csv")):
        return "xlsx"
    if "csv" in content_type or path.endswith(".csv"):
        return "csv"
    return "html"


def _ckan_text(result) -> tuple[str, str, str, str, dict]:
    """Render a CKAN action result. Returns (source_type, title, text, resource_url, metadata)."""
    if isinstance(result, dict) and "url" in result and ("package_id" in result or "format" in result):
        meta = {k: result.get(k) for k in CKAN_RESOURCE_FIELDS if result.get(k) not in (None, "")}
        lines = ["CKAN resource metadata:"] + [f"- {k}: {v}" for k, v in meta.items()]
        lines.append("To read the data itself, call fetch_url on the resource url above.")
        return "ckan_resource", str(result.get("name") or ""), "\n".join(lines), str(result.get("url") or ""), meta
    if isinstance(result, dict) and isinstance(result.get("resources"), list):
        org = (result.get("organization") or {}).get("title", "")
        lines = [f"CKAN dataset: {result.get('title', '')}", f"Publisher: {org}",
                 f"Last modified: {result.get('metadata_modified', '')}", f"Description: {result.get('notes', '')}",
                 "Resources:"]
        lines += [f"- {r.get('name', '')} [{r.get('format', '')}] id={r.get('id', '')} url={r.get('url', '')}"
                  for r in result["resources"][:50] if isinstance(r, dict)]
        meta = {"name": result.get("name"), "title": result.get("title"), "organization": org,
                "metadata_modified": result.get("metadata_modified")}
        return "ckan_dataset", str(result.get("title") or ""), "\n".join(lines), "", meta
    if isinstance(result, dict) and isinstance(result.get("records"), list):
        fields = [f.get("id", "") for f in result.get("fields", []) if isinstance(f, dict)]
        records = result["records"][:MAX_TABLE_ROWS]
        header = fields or (list(records[0]) if records and isinstance(records[0], dict) else [])
        rows = [[r.get(h) for h in header] for r in records if isinstance(r, dict)]
        text = f"CKAN datastore records (resource {result.get('resource_id', '')}):\n" + _render_rows(header, rows, result.get("total"))
        return "ckan_datastore", "", text, "", {"resource_id": result.get("resource_id"), "total": result.get("total")}
    return "json", "", json.dumps(result, ensure_ascii=False, indent=1), "", {}


def extract_json(data: bytes) -> tuple[str, str, str, str, dict]:
    """Parse JSON; CKAN action API responses get a readable rendering.
    Returns (source_type, title, text, resource_url, metadata). Raises ValueError."""
    try:
        doc = json.loads(decode_text(data))
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid JSON: {exc}") from exc
    if isinstance(doc, dict) and "success" in doc and ("result" in doc or "error" in doc):
        if not doc.get("success"):
            err = doc.get("error") or {}
            msg = err.get("message") if isinstance(err, dict) else str(err)
            raise ValueError(f"CKAN API error: {msg or 'request failed'}")
        return _ckan_text(doc.get("result"))
    return "json", "", json.dumps(doc, ensure_ascii=False, indent=1), "", {}


def _download(
    url: str, session: requests.Session, max_bytes: int | None = None, headers: dict | None = None
) -> tuple[requests.Response, bytes]:
    """GET with manual redirect handling (each hop is checked) and a size cap."""
    max_bytes = max_bytes or MAX_DOWNLOAD_BYTES
    current = url
    request_headers = {"User-Agent": USER_AGENT, "Accept-Language": "he,en;q=0.8", **(headers or {})}
    for _ in range(MAX_REDIRECTS + 1):
        problem = _check_url(current)
        if problem:
            raise ValueError(problem)
        resp = session.get(current, headers=request_headers, timeout=TIMEOUT, stream=True, allow_redirects=False)
        if resp.is_redirect and resp.headers.get("Location"):
            current = urljoin(current, resp.headers["Location"])
            resp.close()
            continue
        chunks, total = [], 0
        for chunk in resp.iter_content(64 * 1024):
            total += len(chunk)
            if total > max_bytes:
                resp.close()
                raise ValueError(f"Download exceeds the {max_bytes:,}-byte size limit.")
            chunks.append(chunk)
        return resp, b"".join(chunks)
    raise ValueError("Too many redirects.")


def _retry_delay(attempt: int, resp: requests.Response | None) -> float:
    delay = BACKOFF_BASE_S * (2 ** attempt)
    retry_after = resp.headers.get("Retry-After") if resp is not None else None
    if retry_after and str(retry_after).strip().isdigit():
        delay = max(delay, float(retry_after))
    return min(delay, MAX_BACKOFF_S)


def download_with_retries(
    url: str,
    session: requests.Session,
    max_bytes: int | None = None,
    headers: dict | None = None,
    max_retries: int | None = None,
) -> tuple[requests.Response, bytes, int]:
    """_download plus retries for transient failures only (429, 5xx, connection errors/timeouts).
    401/403/404 and policy errors (ValueError) are never retried. Returns (response, data, attempts)."""
    max_retries = MAX_RETRIES if max_retries is None else max_retries
    attempt = 0
    while True:
        try:
            resp, data = _download(url, session, max_bytes=max_bytes, headers=headers)
        except (requests.ConnectionError, requests.Timeout):
            if attempt >= max_retries:
                raise
            _sleep(_retry_delay(attempt, None))
            attempt += 1
            continue
        if resp.status_code in TRANSIENT_STATUSES and attempt < max_retries:
            _sleep(_retry_delay(attempt, resp))
            attempt += 1
            continue
        return resp, data, attempt + 1


def robots_allowed(url: str, session: requests.Session) -> bool:
    """Respect robots.txt for web pages. Only an explicit Disallow blocks; a missing or unreadable
    robots.txt allows the request (the request itself then decides). Cached per host."""
    parsed = urlparse(url)
    origin = f"{parsed.scheme}://{parsed.netloc}"
    if origin not in _robots_cache:
        parser = None
        try:
            resp, data = _download(origin + "/robots.txt", session, max_bytes=MAX_ROBOTS_BYTES)
            if resp.status_code == 200:
                parser = RobotFileParser()
                parser.parse(decode_text(data).splitlines())
        except (requests.RequestException, ValueError):
            parser = None
        _robots_cache[origin] = parser
    parser = _robots_cache[origin]
    return parser is None or parser.can_fetch(ROBOTS_AGENT, url)


def fetch_url(
    url: str,
    max_chars: int | None = None,
    session: requests.Session | None = None,
    respect_robots: bool = True,
) -> FetchResult:
    """Fetch a URL and return its readable text, page by page (`document`), without truncation unless the
    caller asks for `max_chars`. Never raises: failures are reported in the result."""
    session = session or requests.Session()
    try:
        if respect_robots and _check_url(url) is None and not robots_allowed(url, session):
            return FetchResult(url=url, final_url=url, ok=False, source_type="unknown", access_restricted=True,
                               error="Disallowed by the site's robots.txt; not fetched.")
        resp, data, attempts = download_with_retries(url, session)
    except (requests.RequestException, ValueError) as exc:
        return FetchResult(url=url, final_url=url, ok=False, source_type="unknown", error=str(exc))

    final_url = resp.url or url
    if resp.status_code >= 400:
        error = f"HTTP {resp.status_code}"
        restricted = resp.status_code in (401, 403)
        if restricted:
            error += " (access denied; not retried or bypassed)"
        elif resp.status_code in TRANSIENT_STATUSES:
            error += f" (gave up after {attempts} attempts)"
        return FetchResult(
            url=url, final_url=final_url, ok=False, source_type="unknown",
            http_status=resp.status_code, error=error, attempts=attempts, access_restricted=restricted,
        )

    content_type = (resp.headers.get("Content-Type") or "").lower()
    kind = detect_kind(content_type, final_url, data)
    is_pdf, is_json, is_xlsx, is_csv = (kind == k for k in ("pdf", "json", "xlsx", "csv"))
    resource_url, metadata = "", {}
    source_type = kind
    document = None
    try:
        if is_pdf:
            document = documents.extract_pdf(data)
            title, text = document.title, pages_text(document)
        elif is_json:
            source_type, title, text, resource_url, metadata = extract_json(data)
            if resource_url:
                resource_url = urljoin(final_url, resource_url)
        elif is_xlsx:
            title, text = "", extract_xlsx_text(data)
        elif is_csv:
            title, text = "", extract_csv_text(data)
        elif "html" in content_type or "xml" in content_type or data.lstrip()[:1] == b"<":
            source_type = "html"
            title, text = extract_html_text(data)
        elif content_type.startswith("text/"):
            source_type = "text"
            title, text = "", data.decode(resp.encoding or "utf-8", errors="replace")
        else:
            return FetchResult(
                url=url, final_url=final_url, ok=False, source_type="unknown",
                http_status=resp.status_code, error=f"Unsupported content type: {content_type or 'unknown'}",
            )
    except ValueError as exc:
        return FetchResult(
            url=url, final_url=final_url, ok=False, source_type=source_type,
            http_status=resp.status_code, error=str(exc),
        )
    except Exception as exc:  # parser failure on odd markup or files
        return FetchResult(
            url=url, final_url=final_url, ok=False, source_type=source_type,
            http_status=resp.status_code, error=f"Could not parse content: {exc}",
        )

    if not text.strip():
        return FetchResult(
            url=url, final_url=final_url, ok=False, source_type=source_type,
            http_status=resp.status_code, error="No readable text extracted (page may require JavaScript).",
        )
    text = strip_controls(text)
    if document is None:
        document = documents.extract_sections(text, source_type, title)
        if re.search(r"\[showing first \d+ of \d+ data rows\]", text):
            document.problem("Only the first rows of this table were rendered; use the dataset tools for all records.")
    truncated = False
    if max_chars is not None:
        text, truncated = _truncate(text, max_chars)
    return FetchResult(
        url=url, final_url=final_url, ok=True, source_type=source_type, title=strip_controls(title),
        text=text, http_status=resp.status_code, truncated=truncated,
        resource_url=resource_url, metadata=metadata, attempts=attempts,
        document=document, raw=data if is_pdf else b"",
    )
